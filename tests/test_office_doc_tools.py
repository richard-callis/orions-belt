"""
Tests for the office-document generation MCP tools: create_word_document,
create_powerpoint, create_excel, create_pdf.
"""
import asyncio
import os

import pytest

from app import db
from app.models.connector import AuthorizedDirectory
from app.services.mcp import tools as mcp_tools


def _run(coro):
    return asyncio.get_event_loop().run_until_complete(coro)


@pytest.fixture
def authorized_dir(app, tmp_path):
    with app.app_context():
        d = AuthorizedDirectory(
            path=str(tmp_path), alias="test", recursive=True,
            read_only=False, max_tier=3, enabled=True,
        )
        db.session.add(d)
        db.session.commit()
        yield tmp_path
        AuthorizedDirectory.query.filter_by(path=str(tmp_path)).delete()
        db.session.commit()


class TestCreateWordDocument:
    def test_creates_docx_with_heading_and_bullets(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "report.docx")
            result = _run(mcp_tools._handle_create_word_document("create_word_document", {
                "path": path,
                "title": "Status Report",
                "content": "# Summary\nEverything is on track.\n\n- Item one\n- Item two",
            }))
            assert "Created Word document" in result
            assert os.path.exists(path)

            import docx
            doc = docx.Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "Status Report" in text
            assert "Summary" in text
            assert "Item one" in text

    def test_refuses_outside_authorized_dir(self, app, tmp_path):
        with app.app_context():
            outside = str(tmp_path / "unauthorized.docx")
            result = _run(mcp_tools._handle_create_word_document("create_word_document", {
                "path": outside, "content": "hello",
            }))
            assert result.startswith("Error")
            assert not os.path.exists(outside)

    def test_refuses_if_file_exists(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "exists.docx")
            open(path, "w").close()
            result = _run(mcp_tools._handle_create_word_document("create_word_document", {
                "path": path, "content": "hello",
            }))
            assert "already exists" in result


class TestCreatePowerpoint:
    def test_creates_pptx_with_slides(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "deck.pptx")
            result = _run(mcp_tools._handle_create_powerpoint("create_powerpoint", {
                "path": path,
                "title": "Q1 Review",
                "slides": [
                    {"title": "Wins", "bullets": ["Shipped X", "Fixed Y"]},
                    {"title": "Risks", "bullets": ["Z is slipping"]},
                ],
            }))
            assert "Created PowerPoint" in result
            assert os.path.exists(path)

            from pptx import Presentation
            prs = Presentation(path)
            assert len(prs.slides) == 3  # title slide + 2 content slides

    def test_requires_slides(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "empty.pptx")
            result = _run(mcp_tools._handle_create_powerpoint("create_powerpoint", {"path": path}))
            assert result.startswith("Error")
            assert not os.path.exists(path)


class TestCreateExcel:
    def test_creates_single_sheet(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "data.xlsx")
            result = _run(mcp_tools._handle_create_excel("create_excel", {
                "path": path,
                "sheets": {"headers": ["Name", "Score"], "rows": [["Alice", 90], ["Bob", 85]]},
            }))
            assert "Created Excel workbook" in result

            import openpyxl
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            assert ws["A1"].value == "Name"
            assert ws["A2"].value == "Alice"
            assert ws["B2"].value == 90

    def test_creates_multiple_sheets(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "multi.xlsx")
            result = _run(mcp_tools._handle_create_excel("create_excel", {
                "path": path,
                "sheets": {
                    "Revenue": {"headers": ["Month", "Amount"], "rows": [["Jan", 100]]},
                    "Costs": {"headers": ["Month", "Amount"], "rows": [["Jan", 50]]},
                },
            }))
            assert "2 sheet(s)" in result
            import openpyxl
            wb = openpyxl.load_workbook(path)
            assert set(wb.sheetnames) == {"Revenue", "Costs"}


class TestCreatePdf:
    def test_creates_pdf(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "doc.pdf")
            result = _run(mcp_tools._handle_create_pdf("create_pdf", {
                "path": path,
                "title": "Report",
                "content": "# Overview\nThis is the body.\n\n- point one\n- point two",
            }))
            assert "Created PDF" in result
            assert os.path.exists(path)
            assert os.path.getsize(path) > 0
            with open(path, "rb") as f:
                assert f.read(5) == b"%PDF-"

    def test_refuses_empty_content_and_title(self, app, authorized_dir):
        with app.app_context():
            path = str(authorized_dir / "empty.pdf")
            result = _run(mcp_tools._handle_create_pdf("create_pdf", {"path": path, "content": ""}))
            assert result.startswith("Error")
